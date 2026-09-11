"""Regressionstests: Texte dürfen nicht über den Folienrand hinauslaufen.

Ausführen:  cd autopraesi_python && venv/bin/python -m unittest -v

Hintergrund: Am 13.9.2026 lief ein 215 Zeichen langer Begrüßungsvers bei
54pt über den unteren Rand der Folie. python-pptx kann Text nicht messen,
daher wird die Höhe geschätzt und die Schrift passend verkleinert.
"""
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation  # noqa: E402
from pptx.util import Cm  # noqa: E402

from config import TEMPLATE_PATH  # noqa: E402
from excel_reader import GodiPlanData  # noqa: E402
import presentation_builder as pb  # noqa: E402

VERS_13_9 = ("»Prüft mich hierin, spricht der HERR der Heerscharen, ob ich euch nicht "
             "die Fenster des Himmels auftun werde und Segen herabschütten in Fülle! "
             "[…] Dann werden euch alle Völker glücklich preisen.« (Maleachi 3,10b.12)")

GREETING_SLIDE = 2   # Template-Folie 3
PREDIGT1_SLIDE = 17  # Template-Folie 18


def _slide_with(text_marker, data):
    prs = Presentation(TEMPLATE_PATH)
    for slide in prs.slides:
        if text_marker in pb._get_slide_text(slide):
            pb._fill_slide_simple(slide, data, is_template=True)
            return prs, slide
    raise AssertionError(f"Template-Folie mit {text_marker!r} nicht gefunden")


def _bottom_cm(shape, font_pt):
    """Geschätzte Unterkante (cm) des Textblocks bei gegebener Schriftgröße."""
    return (shape.top + pb._estimate_text_height_pt(shape, font_pt) * pb._EMU_PER_PT) / 360000


class GreetingVerse(unittest.TestCase):
    def test_short_verse_keeps_template_size(self):
        data = GodiPlanData(greeting_verse="Der HERR ist mein Hirte. (Psalm 23,1)")
        _, slide = _slide_with("Begrüßungsvers", data)
        shape = pb._find_shape_with_text(slide, data.greeting_verse)
        self.assertEqual(pb._shape_font_pt(shape), 54)

    def test_long_verse_is_shrunk_to_fit_slide(self):
        data = GodiPlanData(greeting_verse=VERS_13_9)
        prs, slide = _slide_with("Begrüßungsvers", data)
        shape = pb._find_shape_with_text(slide, data.greeting_verse)
        pt = pb._shape_font_pt(shape)
        self.assertLess(pt, 54)
        self.assertGreaterEqual(pt, 24)
        # Textblock endet oberhalb des unteren Folienrands (inkl. 0.8cm Rand)
        self.assertLessEqual(_bottom_cm(shape, pt), prs.slide_height / 360000 - 0.8)
        # … und die nächstgrößere Stufe hätte NICHT mehr gepasst (nicht zu klein)
        self.assertGreater(_bottom_cm(shape, pt + 2), prs.slide_height / 360000 - 0.8)

    def test_estimate_reproduces_observed_overflow(self):
        # Beobachtung aus dem Fehlerbild: ~20 Zeichen je Zeile bei 54pt, Text
        # läuft unten aus der Folie. Das Modell muss diesen Überlauf erkennen.
        prs = Presentation(TEMPLATE_PATH)
        shape = pb._find_shape_with_text(prs.slides[GREETING_SLIDE], "Begrüßungsvers")
        width_pt = (shape.width - 2 * 63281) / pb._EMU_PER_PT
        lines = pb._estimate_line_count(VERS_13_9, width_pt, 54)
        self.assertGreaterEqual(lines, 10)  # 215 Zeichen / ~20 je Zeile
        self.assertLessEqual(lines, 14)
        pb._replace_text_in_shape(shape, "Begrüßungsvers", VERS_13_9)
        self.assertGreater(_bottom_cm(shape, 54), prs.slide_height / 360000)  # = Überlauf

    def test_extremely_long_verse_stops_at_minimum(self):
        data = GodiPlanData(greeting_verse=VERS_13_9 * 4)
        _, slide = _slide_with("Begrüßungsvers", data)
        shape = pb._find_shape_with_text(slide, data.greeting_verse)
        self.assertEqual(pb._shape_font_pt(shape), 24)

    def test_multiline_verse_counts_line_breaks(self):
        single = pb._estimate_line_count("a b c d e", 1000, 20)
        multi = pb._estimate_line_count("a b\nc d\ne", 1000, 20)
        self.assertEqual(single, 1)
        self.assertEqual(multi, 3)

    def test_end_para_size_follows_run_size(self):
        # endParaRPr muss mitziehen, sonst zeigt PowerPoint den Cursor/Leerzeilen zu groß
        from pptx.oxml.ns import qn
        data = GodiPlanData(greeting_verse=VERS_13_9)
        _, slide = _slide_with("Begrüßungsvers", data)
        shape = pb._find_shape_with_text(slide, data.greeting_verse)
        end = shape.text_frame.paragraphs[0]._p.find(qn('a:endParaRPr'))
        self.assertEqual(int(end.get('sz')), int(pb._shape_font_pt(shape) * 100))


class PredigtSlide(unittest.TestCase):
    def test_normal_title_unchanged(self):
        data = GodiPlanData(predigt1_reference="Mal 2,17-3,5",
                            predigt1_title="Wie Gott das bittere Herz erneuert")
        _, slide = _slide_with("Predigt zu Predigtext1", data)
        self.assertEqual(pb._shape_font_pt(pb._find_shape_with_text(slide, data.predigt1_title)), 54)
        self.assertEqual(pb._shape_font_pt(pb._find_shape_with_text(slide, "Predigt zu Mal")), 32)

    def test_long_title_stays_above_reference_line(self):
        title = ("Wie Gott unser müdes, bitteres und ängstliches Herz durch sein Wort "
                 "Schritt für Schritt erneuert, uns neue Hoffnung schenkt und uns "
                 "trotz aller Zweifel wieder in seine Gemeinschaft zurückruft")
        data = GodiPlanData(predigt1_reference="Mal 2,17-3,5", predigt1_title=title)
        _, slide = _slide_with("Predigt zu Predigtext1", data)
        title_shape = pb._find_shape_with_text(slide, title)
        ref_shape = pb._find_shape_with_text(slide, "Predigt zu Mal")
        pt = pb._shape_font_pt(title_shape)
        self.assertLess(pt, 54)
        self.assertLessEqual(_bottom_cm(title_shape, pt), ref_shape.top / 360000 - 0.3)

    def test_long_reference_list_stays_on_slide(self):
        ref = "Mt 6,25–34; 1. Petr 5,5b–11; 1. Mose 2,4b–9(10–14)15(18–25); Ps 23,1-6; Joh 3,16"
        data = GodiPlanData(predigt1_reference=ref, predigt1_title="Thema")
        prs, slide = _slide_with("Predigt zu Predigtext1", data)
        ref_shape = pb._find_shape_with_text(slide, "Predigt zu Mt")
        pt = pb._shape_font_pt(ref_shape)
        self.assertLess(pt, 32)
        self.assertLessEqual(_bottom_cm(ref_shape, pt), prs.slide_height / 360000 - 0.3)


if __name__ == "__main__":
    unittest.main()
