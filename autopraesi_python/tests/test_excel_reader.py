"""Regressionstests für das Einlesen des GoDi-Plans (excel_reader.read_godi_plan_by_sheet).

Ausführen:  cd autopraesi_python && venv/bin/python -m unittest -v

Hintergrund (8.10.2026, GoDi-Plan 2026_Q4): Die Zeilen im Plan hatten sich verschoben
(z.B. "Liedbeitrag" vor der Begrüßung), daher kam als Begrüßungsvers der Liedtitel
"Herr aller Zeit", ein Gemeindelied mit Theme-Farbe ("Sonstige Lieder - …") fehlte,
"Abendmahl" (jetzt Spalte B, Zeile 33) wurde nie erkannt, und die schmaleren
Samstags-Sheets (Inhalte in Spalte C) lieferten Uhrzeiten und Namen als "Lieder".
Die Tests bauen kleine Arbeitsmappen im Speicher nach – ohne Dropbox-Zugang.
"""
import datetime
import io
import os
import sys
import unittest
from unittest import mock

import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import excel_reader  # noqa: E402
from excel_reader import _parse_header, read_godi_plan_by_sheet  # noqa: E402

GREEN = PatternFill("solid", fgColor="FFC5E0B4")
GREEN_SA = PatternFill("solid", fgColor="FF99CC00")
RED = PatternFill("solid", fgColor="FFFFA7A7")
ORANGE = PatternFill("solid", fgColor="FFF8CBAD")
GREY = PatternFill("solid", fgColor="FFF2F2F2")


def _sunday_q4(ws):
    """Ausschnitt von 'So 11.10' aus GoDi-Plan 2026_Q4.xlsx."""
    ws["A1"] = "Gottesdienst am 11.10.2026 (19. Sonntag nach Trinitatis)"
    ws["A3"], ws["D3"] = "Thema:", "Mit Gott durchs Leben – gehalten in jeder Lage!"
    ws["A4"], ws["D4"] = "GoDi-Leitung:", "Lukas Schachtmaier"
    ws["A5"], ws["D5"] = "Predigttextvorschläge:", "Mk 2,1–12 ; Jak 5,13–16"
    ws["A7"], ws["D7"] = "Lesungstext / GoDi-Beginn:", "Psalm 25,1-5"
    ws["B16"], ws["D16"] = "Liedbeitrag", "Herr aller Zeit"
    ws["B17"], ws["D17"] = "Begrüßung", '"Heile du mich, Herr, so werde ich heil." Jeremia 17,14'
    ws["B20"], ws["D20"] = "Gemeindelied", "FJ1 160 - Du bist der Weg und die Wahrheit und das Leben"
    ws["D20"].fill = GREEN
    ws["D22"] = "Lobpreisstrophe: FJ3 161 - Rede du selbst"
    ws["D22"].fill = GREEN
    ws["D25"] = "Kinderlied: Immer und überall"
    ws["D25"].fill = GREEN
    ws["B27"], ws["D27"], ws["E27"] = "Predigt 1 + Gebet", "Psalm 25,5", "Mit Gott leben"
    ws["D27"].fill = RED
    ws["B28"], ws["D28"] = "Liedbeitrag", "Meine Seele ist stille in dir"
    # Gemeindelied OHNE grüne Farbe (Theme-Farbe in der echten Datei)
    ws["B31"], ws["D31"] = "Gemeindelied", "Sonstige Lieder - Gottes guter Segen sei mit dir "
    ws["B34"], ws["D34"], ws["E34"] = ("Predigt 2 + Gebet + VU", "Markus 2, 1-12",
                                       "Mit Gott leben – und getragen werden")
    ws["D34"].fill = RED
    ws["B35"], ws["D35"] = "Lobpreisstrophe", "GLS 5 - Nun danket alle Gott (SUG 91)"
    ws["D35"].fill = GREEN
    ws["B37"] = "Abkündigungen"
    for i, (d, t, e) in enumerate([("Di 13.10.26", datetime.time(19, 0), "Bibelstunde"),
                                   ("Sa 17.10.26", datetime.time(18, 30), "Gottesdienst")]):
        ws.cell(row=38 + i, column=2, value=d)
        ws.cell(row=38 + i, column=3, value=t)
        ws.cell(row=38 + i, column=4, value=e)
    ws["F38"] = "Bitte vormerken!"
    ws["B46"] = "Danke an / Grüße:"
    ws["B47"], ws["D47"] = "Gemeindelied", "SUG 118 - Weiß ich den Weg auch nicht "
    ws["D47"].fill = GREEN
    # Namensliste unten (darf nirgends als Lied/Bibelstelle auftauchen)
    ws["A53"], ws["D53"] = "Alexander Schachtmaier", "Alexander Schachtmaier"


def _saturday_q4(ws):
    """Ausschnitt eines Samstags-Sheets (Inhalte in Spalte C, Kürzel in F)."""
    ws["A1"] = "Gottesdienst am Samstag 03.10."
    ws["A3"], ws["C3"] = "Thema:", "Gottes Führung durch Träume und Visionen"
    ws["A4"], ws["C4"], ws["F4"] = "GoDi-Leitung:", "Viktor Janke", "VJ"
    ws["A5"], ws["C5"] = "Predigttext:", "Joel 3,1-2"
    ws["A6"], ws["C6"] = "Verkündigung:", "Viktor Janke"
    ws["B11"], ws["F11"] = "Begrüßung", "VJ"
    ws["B14"], ws["C14"] = "Gemeindelied", "GLS 45 - Teures Wort aus Gottes Munde"
    ws["C14"].fill = GREEN_SA
    ws["B15"], ws["C15"] = "GoDi-Beginn", "Beginn, Gebet, Lesung"
    ws["B16"], ws["C16"] = "Gemeindelied", "GLS 58 - Du weiser Schöpfer aller Dinge"
    ws["C16"].fill = GREEN_SA
    ws["B17"], ws["C17"] = "Predigt", "Träume und Visionen"
    ws["B19"], ws["C19"] = "Gemeindelied", "GLS 672 - Näher mein Gott zu dir"
    ws["C19"].fill = GREEN_SA
    ws["B20"] = "Abkündigungen, Vater Unser + Segen"
    ws["F20"] = "VJ"
    for i, (d, t, e) in enumerate([("So 04.10.26", datetime.time(10, 30), "Gottesdienst"),
                                   ("Di 06.10.26", datetime.time(19, 0), "Bibelstunde")]):
        ws.cell(row=20 + i, column=3, value=d)
        ws.cell(row=20 + i, column=4, value=t)
        ws.cell(row=20 + i, column=5, value=e)
    ws["B27"], ws["C27"] = "Gemeindelied", "GLS 890 - Sendet Licht"
    ws["C27"].fill = GREEN_SA
    ws["A29"] = "Nach dem GoDi klären:"
    # Namensliste (früher als "Lieder"/Predigt 2 gelesen)
    ws["A33"], ws["D33"], ws["E33"] = "Viktor Janke", "Maxi Grundmann", "Jutta Neumann"
    ws["D34"] = "Markus Reinke"


def _abendmahl_q4(ws):
    """Sonntag mit Abendmahl statt Predigt 2 (Spalte B, Zeile 33) und Platzhaltern."""
    ws["A1"] = "Gottesdienst am 25.10.2026 (21. Sonntag nach Trinitatis)"
    ws["A3"] = "Thema:"
    ws["A4"] = "GoDi-Leitung:"
    ws["A5"] = "Predigttextvorschläge:"
    ws["B16"] = "Begrüßung"
    ws["B26"], ws["D26"], ws["E26"] = "Predigt 1 + Gebet", "Text: ", "Titel: "
    ws["B33"], ws["D33"] = "Abendmahl", "1. Korinther 11,23-26"
    ws["B35"] = "Abkündigungen"
    ws["B37"], ws["C37"], ws["D37"] = "Di 27.10.26", "19:00:00", "Bibelstunde"
    # Zweiter Termin am selben Tag ohne Datum
    ws["C38"], ws["D38"] = datetime.time(20, 0), "Chorprobe"
    ws["B44"] = "Ausblick:"


def _read(builder, sheet="Test", path="/Gemeinde/GoDi-Plan 2026_Q4.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    builder(ws)
    buf = io.BytesIO()
    wb.save(buf)
    with mock.patch.object(excel_reader.storage, "download_bytes", return_value=buf.getvalue()):
        return read_godi_plan_by_sheet(sheet, path)


class SundayQ4Test(unittest.TestCase):
    def setUp(self):
        self.d = _read(_sunday_q4)

    def test_greeting_from_begruessung_row_not_liedbeitrag(self):
        self.assertTrue(self.d.greeting_verse.startswith('"Heile du mich'))

    def test_songs_including_uncoloured_gemeindelied(self):
        songs = {s.slot_key: s.raw for s in self.d.songs if s.raw}
        self.assertEqual(songs["song1"], "FJ1 160 - Du bist der Weg und die Wahrheit und das Leben")
        self.assertEqual(songs["song2"], "Lobpreisstrophe: FJ3 161 - Rede du selbst")
        self.assertEqual(songs["song3"], "Kinderlied: Immer und überall")
        self.assertEqual(songs["song4"], "Sonstige Lieder - Gottes guter Segen sei mit dir")
        self.assertEqual(songs["song6"], "GLS 5 - Nun danket alle Gott (SUG 91)")
        self.assertEqual(songs["song7"], "SUG 118 - Weiß ich den Weg auch nicht")
        self.assertNotIn("song5", songs)
        self.assertFalse(any("Liedbeitrag" in r or "Herr aller Zeit" in r for r in songs.values()))

    def test_references_and_titles(self):
        self.assertEqual(self.d.theme, "Mit Gott durchs Leben – gehalten in jeder Lage!")
        self.assertEqual(self.d.lesung_reference, "Psalm 25,1-5")
        self.assertEqual((self.d.predigt1_reference, self.d.predigt1_title),
                         ("Psalm 25,5", "Mit Gott leben"))
        self.assertEqual(self.d.predigt2_reference, "Markus 2, 1-12")
        self.assertEqual(self.d.layout, "sonntag")
        self.assertFalse(self.d.is_abendmahl)

    def test_invitation_block(self):
        ev = [(e.date_str, e.time_str, e.event_name, e.note) for e in self.d.invitation_events]
        self.assertEqual(ev, [("Di 13.10.26", "19:00", "Bibelstunde", "Bitte vormerken!"),
                              ("Sa 17.10.26", "18:30", "Gottesdienst", "")])


class SaturdayQ4Test(unittest.TestCase):
    def setUp(self):
        self.d = _read(_saturday_q4)

    def test_layout_and_header_without_year(self):
        self.assertEqual(self.d.layout, "samstag")
        self.assertEqual(self.d.date_str, "03.10.2026")
        self.assertEqual(self.d.theme, "Gottes Führung durch Träume und Visionen")

    def test_songs_from_column_c(self):
        songs = {s.slot_key: s.raw for s in self.d.songs if s.raw}
        self.assertEqual(songs, {
            "song1": "GLS 45 - Teures Wort aus Gottes Munde",
            "song4": "GLS 58 - Du weiser Schöpfer aller Dinge",
            "song5": "GLS 672 - Näher mein Gott zu dir",
            "song7": "GLS 890 - Sendet Licht",
        })

    def test_predigt_and_no_name_garbage(self):
        self.assertEqual(self.d.predigt1_reference, "Joel 3,1-2")
        self.assertEqual(self.d.predigt1_title, "Träume und Visionen")
        self.assertEqual((self.d.predigt2_reference, self.d.predigt2_title), ("", ""))
        self.assertEqual(self.d.greeting_verse, "")
        self.assertEqual(self.d.lesung_reference, "")

    def test_invitations_start_in_abkuendigungen_row(self):
        ev = [(e.date_str, e.time_str, e.event_name, e.note) for e in self.d.invitation_events]
        self.assertEqual(ev, [("So 04.10.26", "10:30", "Gottesdienst", ""),
                              ("Di 06.10.26", "19:00", "Bibelstunde", "")])


class AbendmahlQ4Test(unittest.TestCase):
    def setUp(self):
        self.d = _read(_abendmahl_q4)

    def test_abendmahl_detected_in_column_b(self):
        self.assertTrue(self.d.is_abendmahl)

    def test_placeholders_and_abendmahl_row_not_used_as_sermon(self):
        self.assertEqual((self.d.predigt1_reference, self.d.predigt1_title), ("", ""))
        self.assertEqual((self.d.predigt2_reference, self.d.predigt2_title), ("", ""))

    def test_second_event_same_day_inherits_date(self):
        ev = [(e.date_str, e.time_str, e.event_name) for e in self.d.invitation_events]
        self.assertEqual(ev, [("Di 27.10.26", "19:00", "Bibelstunde"),
                              ("Di 27.10.26", "20:00", "Chorprobe")])


class HeaderTest(unittest.TestCase):
    def test_header_variants(self):
        self.assertEqual(_parse_header("Gottesdienst am 08.03.2026 (Okuli)"), ("08.03.2026", "Okuli"))
        self.assertEqual(_parse_header("Gottesdienst am Samstag 03.10.",
                                       "/Gemeinde/GoDi-Plan 2026_Q4.xlsx"), ("03.10.2026", ""))
        # Januar in einem Q4-Plan gehört ins Folgejahr
        self.assertEqual(_parse_header("Gottesdienst am Samstag 2.1.",
                                       "/Gemeinde/GoDi-Plan 2026_Q4.xlsx"), ("02.01.2027", ""))


if __name__ == "__main__":
    unittest.main()
